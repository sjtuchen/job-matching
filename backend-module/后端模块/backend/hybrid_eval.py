# -*- coding: utf-8 -*-
"""
⑥ 混合匹配评测：图谱清单 + BGE 语义层 融合
- 与 graph_match.py 同口径（族级平均分排名、学历门槛、分层 Top1/Top3）
- 同一份图谱缓存同时跑 α=1.0(纯清单)/0.7/0.5/0.3 四档融合权重 → 严格可比消融
- 语义分：双向最大覆盖均值，底噪校准（BGE 不相关对 ~0.3 → 0）
- 负对照隔离：零交集(严格版) 与 语义阈值≥0.6(语义版) 两种模式对比
零 API 调用：全部读本地缓存（graph_cache + bge_sim_matrix.json）
"""
import json
import os
import re
from collections import defaultdict

BASE = r"D:\新国赛"
TEMP = os.path.join(BASE, ".temp")
CACHE_DIR = os.path.join(TEMP, "graph_cache")
LIST_PATH = os.path.join(BASE, r"简历语料\合成样本清单.json")
SIM_PATH = os.path.join(TEMP, "bge_sim_matrix.json")

W_SKILL, W_TOOL, W_SOFT = 0.5, 0.2, 0.3
NEG_FAMS = {"半导体", "先进制造", "智能制造", "生物", "新能源车", "绿色低碳", "嵌入式"}
BASE_NOISE = 0.35      # BGE 底噪：低于此视为不相关
NEG_SIM_GATE = 0.60    # 负对照语义隔离阈值


def load_graphs():
    graphs = {}
    for f in os.listdir(CACHE_DIR):
        if f.endswith(".json"):
            with open(os.path.join(CACHE_DIR, f), encoding="utf-8") as fh:
                g = json.load(fh)
                graphs[g["doc_id"]] = g
    return graphs


def load_jd_meta():
    from openpyxl import load_workbook
    meta = {}
    for path in [os.path.join(BASE, "jd_data.xlsx"),
                 os.path.join(BASE, r"简历语料\boss_AIGC补充24条.xlsx")]:
        wb = load_workbook(path, read_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        header_idx = next(i for i, r in enumerate(rows) if r and r[0] == "jd_id")
        col = {h: i for i, h in enumerate(rows[header_idx])}
        for row in rows[header_idx + 1:]:
            if not row or not row[0]:
                continue
            meta[str(row[col["jd_id"]])] = {
                "族": (row[col["岗位类别"]] or "").strip(),
                "岗位名称": row[col["岗位名称"]],
                "学历": (row[col["学历要求"]] or "").strip(),
            }
    return meta


def tags(g):
    sg = g["skill_groups"]
    return (set(sg["专业技能"]), set(sg["工具"]), set(sg["软技能"]))


def edu_gate(r_edu, jd_edu):
    order = {"大专": 0, "本科": 1, "硕士": 2, "博士": 3}
    m = re.search(r"(大专|本科|硕士|博士)", jd_edu or "")
    if not m:
        return True
    return order.get(r_edu, 1) >= order[m.group(1)]


def jaccard(a, b):
    if not a and not b:
        return 1.0
    if not a or not b:
        return 0.0
    return len(a & b) / len(a | b)


# ---------- 语义层 ----------
class Sem:
    def __init__(self):
        with open(SIM_PATH, encoding="utf-8") as f:
            self.sim = json.load(f)["sim"]

    def pair(self, a, b):
        # 矩阵为上三角存储，必须双向查询
        va = self.sim.get(a, {}).get(b)
        if va is None:
            va = self.sim.get(b, {}).get(a, 0.0)
        return va

    def cover(self, need, have):
        """need 中每个词在 have 中的最大相似度均值（需求覆盖度）"""
        if not need or not have:
            return 0.0
        vals = [max((self.pair(a, b) for b in have), default=BASE_NOISE) for a in need]
        return sum(vals) / len(vals)

    def align(self, ra, rb):
        """双向对齐分（0~1，底噪已校准）"""
        if not ra or not rb:
            return 0.0
        def cal(v):
            return max(0.0, (v - BASE_NOISE) / (1.0 - BASE_NOISE))
        return 0.5 * cal(self.cover(rb, ra)) + 0.5 * cal(self.cover(ra, rb))

    def has_sem_overlap(self, ra, rb):
        """是否存在一对技能词语义相似度超阈值（负对照语义隔离用）"""
        return any(self.pair(a, b) >= NEG_SIM_GATE for a in ra for b in rb)


def hybrid_score(rg, jg, sem, alpha):
    """alpha=1.0 纯清单；alpha<1 时 (alpha*清单 + (1-alpha)*语义) 逐组融合"""
    if not edu_gate(rg.get("education", "本科"), jg.get("hard_requirements", {}).get("学历", "")):
        return None  # 学历门槛出局
    r_skill, r_tool, r_soft = tags(rg)
    j_skill, j_tool, j_soft = tags(jg)

    parts = []
    for r_set, j_set, w in [(r_skill, j_skill, W_SKILL),
                            (r_tool, j_tool, W_TOOL),
                            (r_soft, j_soft, W_SOFT)]:
        s = jaccard(r_set, j_set)
        if alpha < 1.0:
            s = alpha * s + (1 - alpha) * sem.align(r_set, j_set)
        parts.append(w * s)
    return sum(parts)


def eval_alpha(r_graphs, jd_graphs, jd_meta, samples, sem, alpha, neg_mode):
    """跑一档配置，返回 {层: {rank: cnt}} 和逐份明细"""
    rank_by_layer = defaultdict(lambda: defaultdict(int))
    details = []
    for rid, rg in sorted(r_graphs.items()):
        s = samples.get(rid, {})
        fam_scores, cnt = defaultdict(float), defaultdict(int)
        for jd_id, jg in jd_graphs.items():
            fam = jd_meta.get(jd_id, {}).get("族", "")
            # 负对照隔离
            if any(nf in fam for nf in NEG_FAMS):
                r_skill, _, _ = tags(rg)
                j_skill, _, _ = tags(jg)
                if neg_mode == "strict":
                    if not (r_skill & j_skill):
                        continue
                else:  # semantic
                    if not sem.has_sem_overlap(r_skill, j_skill):
                        continue
            sc = hybrid_score(rg, jg, sem, alpha)
            if sc is None:
                continue
            fam_scores[fam] += sc
            cnt[fam] += 1
        if not cnt:
            continue
        fam_med = {f: fam_scores[f] / cnt[f] for f in cnt}
        target = s.get("目标族", "")
        target_score = fam_med.get(target)
        if target_score is None:
            rank_by_layer[s.get("层", "?")][99] += 1  # 目标族被过滤，记为垫底
            details.append({"id": rid, "层": s.get("层"), "目标族": target, "rank": None})
            continue
        rank = sum(1 for f, v in fam_med.items() if v > target_score) + 1
        rank_by_layer[s.get("层", "?")][rank] += 1
        details.append({"id": rid, "层": s.get("层"), "目标族": target, "rank": rank})
    return rank_by_layer, details


def fmt_layer(c):
    n = sum(c.values())
    top1 = c[1]
    top3 = sum(v for k, v in c.items() if k <= 3)
    return n, top1, top3


def main():
    graphs = load_graphs()
    jd_graphs = {k: v for k, v in graphs.items() if v.get("doc_type") == "jd"}
    r_graphs = {k: v for k, v in graphs.items() if v.get("doc_type") == "resume"}
    print(f"图谱库：JD {len(jd_graphs)} 份，简历 {len(r_graphs)} 份")
    if not jd_graphs or not r_graphs:
        print("图谱库未就绪")
        return
    jd_meta = load_jd_meta()
    with open(LIST_PATH, encoding="utf-8") as f:
        samples = {s["id"]: s for s in json.load(f)}
    sem = Sem()

    configs = []
    for neg_mode in ["strict", "semantic"]:
        for alpha in [1.0, 0.7, 0.5, 0.3]:
            configs.append((neg_mode, alpha))

    all_results = {}
    print()
    print("=" * 78)
    print(f"⑥ 混合匹配评测（JD {len(jd_graphs)} 份缓存版；基线为 359 全量口径，仅作参考）")
    print("=" * 78)
    hdr = f"{'负对照':<10}{'α清单':<8}{'对口Top1/Top3':<18}{'迁移Top1/Top3':<18}{'跨界Top1/Top3'}"
    print(hdr)
    report_rows = []
    for neg_mode, alpha in configs:
        rank_by_layer, details = eval_alpha(r_graphs, jd_graphs, jd_meta, samples, sem, alpha, neg_mode)
        all_results[(neg_mode, alpha)] = (rank_by_layer, details)
        cells = []
        for layer in ["对口层", "迁移层", "跨界层"]:
            n, t1, t3 = fmt_layer(rank_by_layer[layer])
            cells.append(f"{t1}/{n} {t3}/{n}")
        label_a = f"{alpha:.1f}" if alpha < 1 else "1.0纯清单"
        print(f"{neg_mode:<10}{label_a:<8}{cells[0]:<18}{cells[1]:<18}{cells[2]}")
        report_rows.append({
            "负对照模式": neg_mode, "alpha": alpha,
            "分层": {layer: {"n": fmt_layer(rank_by_layer[layer])[0],
                            "Top1": fmt_layer(rank_by_layer[layer])[1],
                            "Top3": fmt_layer(rank_by_layer[layer])[2]}
                    for layer in ["对口层", "迁移层", "跨界层"]},
        })

    # 选最优配置存逐份明细
    def score_config(row):
        # 主指标：三层 Top1 之和（跨界层权重×2，最难）
        d = row["分层"]
        return d["对口层"]["Top1"] + d["迁移层"]["Top1"] + 2 * d["跨界层"]["Top1"]
    best_row = max(report_rows, key=score_config)
    best_key = (best_row["负对照模式"], best_row["alpha"])
    print()
    print(f"最优配置：负对照={best_key[0]}，α={best_key[1]} "
          f"(对口 {best_row['分层']['对口层']['Top1']}/{best_row['分层']['对口层']['n']}，"
          f"迁移 {best_row['分层']['迁移层']['Top1']}/{best_row['分层']['迁移层']['n']}，"
          f"跨界 {best_row['分层']['跨界层']['Top1']}/{best_row['分层']['跨界层']['n']})")

    rank_by_layer, details = all_results[best_key]
    out = {
        "_meta": {"时间": "2026-09-06",
                  "JD图谱数": len(jd_graphs),
                  "说明": "清单Jaccard(0.5/0.2/0.3) 与 BGE 语义层逐组融合；"
                          "语义分=双向最大覆盖均值+底噪0.35校准；负对照语义阈值0.6",
                  "消融": report_rows,
                  "最优配置": {"负对照模式": best_key[0], "alpha": best_key[1]}},
        "最优配置_分层排名分布": {layer: {f"第{k}名": v for k, v in sorted(c.items())}
                                for layer, c in rank_by_layer.items()},
        "最优配置_逐份明细": details,
    }
    out_path = os.path.join(TEMP, "hybrid_eval_report.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
    print(f"报告已存 {out_path}")


if __name__ == "__main__":
    main()
