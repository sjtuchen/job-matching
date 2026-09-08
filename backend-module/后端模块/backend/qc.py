# -*- coding: utf-8 -*-
"""P4 质检：总量/去重/完整度/字段分布/异常值扫描，结果写文件"""
from openpyxl import load_workbook
import re
import json

DATA = r"D:\新国赛\jd_data.xlsx"
OUT = r"D:\新国赛\.temp\qc_report.txt"

wb = load_workbook(DATA)
ws = wb["JD数据"]

rows = []
for r in range(2, ws.max_row + 1):
    v = ws.cell(row=r, column=1).value
    if v in (None, "") or str(v).startswith("AIGC:"):
        continue
    rows.append({
        "row": r,
        "id": str(v),
        "title": str(ws.cell(row=r, column=2).value or ""),
        "company": str(ws.cell(row=r, column=3).value or ""),
        "industry": str(ws.cell(row=r, column=4).value or ""),
        "ent": str(ws.cell(row=r, column=5).value or ""),
        "cat": str(ws.cell(row=r, column=6).value or ""),
        "city_sal": str(ws.cell(row=r, column=7).value or ""),
        "edu": str(ws.cell(row=r, column=8).value or ""),
        "jd": str(ws.cell(row=r, column=9).value or ""),
        "platform": str(ws.cell(row=r, column=10).value or ""),
        "link": str(ws.cell(row=r, column=11).value or ""),
    })

lines = []
lines.append("总条数: %d" % len(rows))

# 1. 链接去重检查
links = [x["link"] for x in rows]
dup_links = len(links) - len(set(links))
lines.append("重复链接: %d" % dup_links)

# 2. 同名同公司重复
pairs = set()
dup_pairs = 0
seen = set()
for x in rows:
    key = (x["title"], x["company"])
    if key in seen:
        dup_pairs += 1
    seen.add(key)

lines.append("同名同公司重复: %d" % dup_pairs)

# 3. 字段完整度
missing = {"title": 0, "company": 0, "industry": 0, "ent": 0, "cat": 0, "edu": 0, "jd": 0, "link": 0}
for x in rows:
    for k in missing:
        if not x[k]:
            missing[k] += 1
lines.append("字段缺失: " + json.dumps(missing, ensure_ascii=False))

# 4. JD 长度分布
lens = [len(x["jd"]) for x in rows]
lens.sort()
if lens:
    lines.append("JD长度 min/中位/max: %d/%d/%d" % (lens[0], lens[len(lens)//2], lens[-1]))
    short = sum(1 for L in lens if L < 150)
    lines.append("JD<150字的条数: %d" % short)

# 5. 类别分布
cats = {}
for x in rows:
    cats[x["cat"]] = cats.get(x["cat"], 0) + 1
lines.append("类别分布: " + json.dumps(cats, ensure_ascii=False))

# 6. 企业性质分布
ents = {}
for x in rows:
    ents[x["ent"]] = ents.get(x["ent"], 0) + 1
lines.append("企业性质: " + json.dumps(ents, ensure_ascii=False))

# 7. 学历分布
edus = {}
for x in rows:
    edus[x["edu"] or "缺失"] = edus.get(x["edu"] or "缺失", 0) + 1
lines.append("学历分布: " + json.dumps(edus, ensure_ascii=False))

# 8. 公司分布 TOP10（检查刷屏）
comps = {}
for x in rows:
    comps[x["company"]] = comps.get(x["company"], 0) + 1
top = sorted(comps.items(), key=lambda t: -t[1])[:10]
lines.append("公司TOP10: " + json.dumps([("%s x%d" % (c, n)) for c, n in top], ensure_ascii=False))

# 9. 异常值扫描：城市字段里的异常、薪资里的异常
weird_city = [x["city_sal"] for x in rows if not x["city_sal"] or "未标注" in x["city_sal"]]
lines.append("地点未标注条数: %d" % len(weird_city))
weird_sal = [x["city_sal"] for x in rows if re.search(r"\d{5,}", x["city_sal"] or "")]
lines.append("薪资疑似异常(5位数字): %d 条，样例: %s" % (len(weird_sal), weird_sal[:3]))

with open(OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("done")
