# -*- coding: utf-8 -*-
"""P4 清理：删9条完全重复JD、从原文补学历、行业补未标注、顺序重编号"""
from openpyxl import load_workbook
import hashlib
import re

DATA = r"D:\新国赛\jd_data.xlsx"

wb = load_workbook(DATA)
ws = wb["JD数据"]

# 收集数据行
rows = []
for r in range(2, ws.max_row + 1):
    v = ws.cell(row=r, column=1).value
    if v in (None, "") or str(v).startswith("AIGC:"):
        continue
    rows.append(r)

# 1. 找出 (公司, JD文本md5) 重复的行，保留首次出现
seen = {}
to_delete = []
for r in rows:
    comp = str(ws.cell(row=r, column=3).value or "")
    jd = str(ws.cell(row=r, column=9).value or "")
    key = (comp, hashlib.md5(jd.encode()).hexdigest())
    if key in seen:
        to_delete.append(r)
        print("删重复: 行%d %s @ %s" % (r, str(ws.cell(row=r, column=2).value)[:20], comp[:18]))
    else:
        seen[key] = r

for r in sorted(to_delete, reverse=True):
    ws.delete_rows(r)
print("共删除 %d 条完全重复" % len(to_delete))

# 2. 补学历（从JD原文）与行业
edu_pat = re.compile(r"(博士|硕士|大专|本科)(及以上)?(学历)?")
fixed_edu, fixed_ind = 0, 0
for r in range(2, ws.max_row + 1):
    v = ws.cell(row=r, column=1).value
    if v in (None, "") or str(v).startswith("AIGC:"):
        continue
    edu = str(ws.cell(row=r, column=8).value or "").strip()
    if not edu:
        jd = str(ws.cell(row=r, column=9).value or "")
        m = edu_pat.search(jd[:600])
        if m:
            ws.cell(row=r, column=8, value=m.group(1))
            fixed_edu += 1
        else:
            ws.cell(row=r, column=8, value="未标注")
    if not str(ws.cell(row=r, column=4).value or "").strip():
        ws.cell(row=r, column=4, value="未标注")
        fixed_ind += 1
print("补学历: %d 条（从JD原文提取），其余记未标注；补行业: %d 条" % (fixed_edu, fixed_ind))

# 3. 顺序重编号
n = 0
for r in range(2, ws.max_row + 1):
    v = ws.cell(row=r, column=1).value
    if v in (None, "") or str(v).startswith("AIGC:"):
        continue
    n += 1
    ws.cell(row=r, column=1, value="JD-%03d" % n)
print("重编号完成: JD-001 ~ JD-%03d" % n)

wb.save(DATA)
print("总条数:", n)
