# -*- coding: utf-8 -*-
"""
JD 半自动采集器（国聘 + 应届生求职网）
用法：
  1. 双击启动：脚本打开浏览器，你正常浏览国聘/应届生的职位详情页
  2. 在详情页上按 Ctrl+Shift+S（或点页面右下角悬浮按钮）
  3. 弹窗确认岗位类别 → 自动提取全部字段 → 追加写入 Excel
  4. 重复浏览-按键即可，最后关窗口自动保存
数据文件：D:/新国赛/jd_data.xlsx（首次运行自动从模板创建）
"""
import datetime
import os
import re
import sys
import threading

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

DATA_PATH = r"D:\新国赛\jd_data.xlsx"
TEMPLATE_PATH = r"D:\新国赛\jd_采集模板.xlsx"

CATEGORIES = [
    "AI/AIGC内容与运营", "数据分析", "数字营销与用户增长", "跨境电商与外贸数字化",
    "新质管理与服务", "新能源/环保文科岗", "算法/研发（对照）", "智能制造/硬件（对照）",
]

# ---------- Excel 读写 ----------

def ensure_data_file():
    if not os.path.exists(DATA_PATH):
        if os.path.exists(TEMPLATE_PATH):
            import shutil
            shutil.copyfile(TEMPLATE_PATH, DATA_PATH)
            # 只保留表头：清掉模板里的示例行和系统水印行
            wb = load_workbook(DATA_PATH)
            ws = wb["JD数据"]
            if ws.max_row > 1:
                ws.delete_rows(2, ws.max_row - 1)
            wb.save(DATA_PATH)
            print("已从模板创建", DATA_PATH)
        else:
            raise SystemExit("找不到模板，请先运行 make_template.py")

def append_row(row):
    wb = load_workbook(DATA_PATH)
    ws = wb["JD数据"]
    # 找第一个可写行：空行，或系统自动注入的 AIGC 水印行（以 AIGC: 开头，直接清掉覆盖）
    r = 2
    while r <= ws.max_row:
        v = ws.cell(row=r, column=1).value
        if v in (None, ""):
            break
        if str(v).startswith("AIGC:"):
            for c in range(1, ws.max_column + 1):
                ws.cell(row=r, column=c).value = None
            print("已清除系统水印行:", r)
            break
        r += 1
    # 去重：同链接已存在则跳过
    link = row[10] or ""
    for i in range(2, r):
        if (ws.cell(row=i, column=11).value or "") == link and link:
            print("跳过重复链接:", link)
            return False
    for col, v in enumerate(row, 1):
        ws.cell(row=r, column=col, value=v)
    wb.save(DATA_PATH)
    print(f"已写入第 {r} 行:", row[1], "|", row[2])
    return True

def next_id():
    wb = load_workbook(DATA_PATH)
    ws = wb["JD数据"]
    r, maxnum = 2, 0
    while ws.cell(row=r, column=1).value not in (None, ""):
        v = ws.cell(row=r, column=1).value
        m = re.match(r"JD-(\d+)", str(v))
        if m:
            maxnum = max(maxnum, int(m.group(1)))
        r += 1
    return f"JD-{maxnum + 1:03d}"

def stats():
    wb = load_workbook(DATA_PATH)
    ws = wb["JD数据"]
    counts = {}
    r = 2
    while ws.cell(row=r, column=1).value not in (None, ""):
        cat = ws.cell(row=r, column=6).value
        if cat:
            counts[cat] = counts.get(cat, 0) + 1
        r += 1
    return counts

# ---------- 国聘页面提取 ----------

JS_IGUOPIN = """
() => {
  const result = {};
  const bodyText = document.body.innerText;

  // 标题：详情页主内容区第一个 h1/h2，或从 title 标签取（格式：职位名-公司名-国聘）
  let title = '';
  const main = document.querySelector('.job-detail, .detail, main, [class*=detail]');
  if (main) {
    const h = main.querySelector('h1, h2, .job-name, [class*=job-name]');
    if (h) title = h.textContent.trim();
  }
  if (!title) title = document.title.split('-')[0].trim();
  result.title = title;

  // 薪资：主内容区中独立的大号"面议"或"xx-xxK"
  const salMatch = bodyText.match(/(\\d+\\s*-\\s*\\d+\\s*[Kk]|\\d+[Kk]以上|面议)\\s*\\n/);
  result.salary = salMatch ? salMatch[1].replace(/\\s/g, '') : '面议';

  // 学历/经验：标签格式"最低学历：xxx"
  const eduMatch = bodyText.match(/最低学历：\\s*([^\\s\\n]+)/);
  result.edu = eduMatch ? eduMatch[1] : '';

  // 地点：工作地点区块
  const locMatch = bodyText.match(/工作地点\\n([^\\n]+)/);
  result.location = locMatch ? locMatch[1].trim() : '';

  // JD 正文：找含"工作职责"的最小文本块
  let jd = '';
  let minLen = Infinity;
  for (const el of document.querySelectorAll('div,section,article')) {
    const t = el.textContent || '';
    if ((t.includes('工作职责') || t.includes('岗位职责') || t.includes('职位描述') || t.includes('任职要求')) && t.length > 100) {
      if (t.length < minLen) { jd = t; minLen = t.length; }
    }
  }
  const idx = jd.search(/工作职责|岗位职责|职位描述|任职要求/);
  result.jd = idx >= 0 ? jd.slice(idx).trim() : jd.slice(0, 3000).trim();

  // 公司：侧栏"单位信息"板块的公司链接（排除图片链接）
  let company = '';
  const companyLinks = [...document.querySelectorAll('a[href*="/company"]')];
  for (const a of companyLinks) {
    const name = a.textContent.trim();
    if (name.length > 3 && name !== '友好企业' && !name.includes('个在招职位')) {
      company = name; break;
    }
  }
  result.company = company;

  // 行业/性质：单位信息下方的元信息行
  // 实际格式（innerText 无分隔符）：用人单位自主招聘国企软件和信息技术服务业30000人以上
  // 结构：{招聘方式}{企业性质}{行业}{规模}人以上（必须用正则字面量，不用 new RegExp 字符串拼接，避免转义坑）
  let ent_type = '', industry = '';
  const m = bodyText.match(/用人单位自主招聘(外商独资|中外合资|事业单位|股份制企业|上市公司|民营企业|央企|国企|民企|外企|合资)([^\\d\\n]{2,30}?)(\\d+[-~\\d]*人[以内以上]?)/);
  if (m) {
    ent_type = m[1].trim();
    industry = m[2].trim();
  } else {
    // 备选：不带"用人单位自主招聘"前缀
    const m2 = bodyText.match(/(国企|民企|民营企业|外企|外商独资|事业单位|上市公司|股份制企业|合资)([^\\d\\n]{2,30}?)(\\d+[-~\\d]*人[以内以上]?)/);
    if (m2) { ent_type = m2[1].trim(); industry = m2[2].trim(); }
  }
  result.ent_type = ent_type;
  result.industry = industry;
  result.url = location.href;
  return result;
}
"""

# ---------- 应届生求职网提取 ----------

JS_YJS = """
() => {
  const result = {};
  const bodyText = document.body.innerText;
  result.title = document.title.split('-')[0].trim() || (document.querySelector('h1') ? document.querySelector('h1').textContent.trim() : '');
  result.company = '';
  const companyMatch = bodyText.match(/(招聘|聘)\\s*[:：]?\\s*([\\u4e00-\\u9fa5（）()]+?(?:公司|集团|银行|大学|学院|研究院|所))$/m);
  result.salary = (bodyText.match(/(\\d+-\\d+K\\/?[\\u4e00-\\u9fa5]*|面议)/) || [])[1] || '';
  result.edu = (bodyText.match(/(本科及以上|硕士及以上|大专及以上|学历不限|本科|硕士|博士)/) || [])[1] || '';
  result.location = (bodyText.match(/工作地点[:：]?\\s*([^\\s，。]+)/) || [])[1] || '';
  const main = document.querySelector('.job-detail, .job-info, article, .content');
  let jd = main ? main.innerText : bodyText;
  const idx = jd.search(/岗位职责|工作职责|职位描述|任职要求|招聘要求/);
  result.jd = idx >= 0 ? jd.slice(idx).slice(0, 4000).trim() : jd.slice(0, 4000).trim();
  result.ent_type = '';
  result.industry = '';
  result.url = location.href;
  return result;
}
"""

# ---------- 注入的悬浮按钮 ----------

INJECT_JS = """
(cid) => {
  if (document.getElementById('jd-collect-btn')) return;
  const btn = document.createElement('button');
  btn.id = 'jd-collect-btn';
  btn.textContent = '📌采集此JD';
  Object.assign(btn.style, {
    position: 'fixed', right: '20px', bottom: '20px', zIndex: 99999,
    padding: '12px 20px', background: '#2F5597', color: '#fff',
    border: 'none', borderRadius: '8px', fontSize: '14px',
    cursor: 'pointer', boxShadow: '0 2px 8px rgba(0,0,0,.3)'
  });
  btn.onclick = () => window.__collect(cid);
  document.body.appendChild(btn);
  console.log('[采集器] 悬浮按钮已注入');
}
"""

SELECT_JS = """
(categories) => new Promise(resolve => {
  // 简易选择弹层
  const old = document.getElementById('jd-cat-overlay');
  if (old) old.remove();
  const overlay = document.createElement('div');
  overlay.id = 'jd-cat-overlay';
  Object.assign(overlay.style, {
    position: 'fixed', inset: '0', background: 'rgba(0,0,0,.5)',
    zIndex: 100000, display: 'flex', alignItems: 'center', justifyContent: 'center'
  });
  const box = document.createElement('div');
  Object.assign(box.style, {
    background: '#fff', borderRadius: '10px', padding: '24px',
    minWidth: '320px', fontFamily: 'sans-serif'
  });
  box.innerHTML = '<div style="font-weight:bold;margin-bottom:12px;">选择岗位类别</div>';
  categories.forEach((c, i) => {
    const b = document.createElement('button');
    b.textContent = c;
    Object.assign(b.style, {
      display: 'block', width: '100%', margin: '6px 0', padding: '10px',
      border: '1px solid #d0d0d0', borderRadius: '6px', background: '#f7f9fc',
      cursor: 'pointer', fontSize: '14px', textAlign: 'left'
    });
    b.onmouseenter = () => b.style.background = '#e3ecff';
    b.onmouseleave = () => b.style.background = '#f7f9fc';
    b.onclick = () => { overlay.remove(); resolve(i); };
    box.appendChild(b);
  });
  overlay.appendChild(box);
  document.body.appendChild(overlay);
  overlay.onclick = e => { if (e.target === overlay) { overlay.remove(); resolve(-1); } };
});
"""

def normalize_ent(t):
    if not t: return "其他"
    if "国企" in t or "央企" in t or "国有" in t: return "国企央企"
    if "民企" in t or "民营" in t: return "民企"
    if "外" in t: return "外企"
    if "事业" in t: return "事业单位"
    if "上市" in t or "股份" in t: return "民企"
    return "其他"

def main():
    from playwright.sync_api import sync_playwright

    ensure_data_file()

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False, args=["--start-maximized"])
        ctx = browser.new_context(no_viewport=True, locale="zh-CN")
        page = ctx.new_page()
        page.goto("https://www.iguopin.com/job/list?keyword=数据分析")
        print("浏览器已打开。浏览到职位详情页后，点右下角【采集此JD】按钮或按 Ctrl+Shift+S。")
        print("关闭浏览器窗口即结束采集（数据已实时保存）。")

        def do_collect():
            try:
                url = page.url
                print("\n[采集] 当前页面:", url)
                if "iguopin.com/job/detail" in url:
                    data = page.evaluate(JS_IGUOPIN)
                    platform = "国聘"
                elif "yingjiesheng.com" in url:
                    data = page.evaluate(JS_YJS)
                    platform = "应届生求职网"
                else:
                    print("[跳过] 不是支持的详情页（国聘详情页 / 应届生网）")
                    return
                print("[提取]", data.get("title"), "|", data.get("company"))
                # 选类别
                idx = page.evaluate(SELECT_JS, CATEGORIES)
                if idx < 0:
                    print("[取消] 未选择类别")
                    return
                cat = CATEGORIES[idx]
                jd_text = (data.get("jd") or "").strip()
                if len(jd_text) < 50:
                    print("[警告] JD 正文过短（%d 字），仍会写入，请人工核对" % len(jd_text))
                row = [
                    next_id(),
                    (data.get("title") or "").strip()[:60],
                    (data.get("company") or "").strip()[:60],
                    (data.get("industry") or "").strip()[:40],
                    normalize_ent(data.get("ent_type")),
                    cat,
                    f"{(data.get('location') or '未标').strip()[:20]}·{data.get('salary') or '面议'}",
                    (data.get("edu") or "").strip()[:10],
                    jd_text,
                    platform,
                    url,
                    datetime.date.today().isoformat(),
                ]
                ok = append_row(row)
                if ok:
                    page.evaluate("() => { const b=document.getElementById('jd-collect-btn'); if(b){b.textContent='✅ 已入库 '+new Date().toLocaleTimeString(); setTimeout(()=>b.textContent='📌采集此JD',2000);} }")
                    print("[完成]", row[0], "当前统计:", stats())
            except Exception as e:
                print("[错误]", repr(e))

        # 快捷键 Ctrl+Shift+S
        page.keyboard  # noqa

        def on_key(event):
            if event.get("key") in ("s", "S") and event.get("ctrlKey") and event.get("shiftKey"):
                do_collect()

        # 给每个新页面注入悬浮按钮 + 快捷键
        def setup_page(pg):
            pg.expose_function("__collect", lambda: do_collect())
            pg.add_init_script("""
                document.addEventListener('keydown', e => {
                    if (e.ctrlKey && e.shiftKey && (e.key === 's' || e.key === 'S')) {
                        e.preventDefault();
                        window.__collect && window.__collect();
                    }
                });
            """)
            pg.on("load", lambda: pg.evaluate(INJECT_JS))
            # 首页立即注入
            pg.evaluate(INJECT_JS)

        ctx.on("page", setup_page)
        setup_page(page)

        # 主线程等待浏览器关闭
        try:
            while True:
                page.wait_for_timeout(2000)
        except Exception:
            pass
        finally:
            print("\n=== 最终统计 ===")
            for k, v in stats().items():
                print(f"  {k}: {v}")
            print("数据文件:", DATA_PATH)
            try:
                browser.close()
            except Exception:
                pass

if __name__ == "__main__":
    main()
